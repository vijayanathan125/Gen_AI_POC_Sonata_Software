import logging
from flask import Flask, request, jsonify
import io
from flask_cors import CORS
import openai
from azure.storage.blob import BlobServiceClient
import docx
import pandas as pd
import pptx
import PyPDF2
#import odf.opendocument
#import odf.text

# Set up logging
logging.basicConfig(filename='error.log', level=logging.ERROR)

# Set your Azure Storage account connection string
azure_storage_connection_string = "AccountName=devstoreaccount1;AccountKey=Eby8vdM02xNOcqFlqUwJPLlmEtlCDXJ1OUzFT50uSRZ6IFsuFq2UVErCz4I6tq/K1SZFPTOtr/KBHBeksoGMGw==;DefaultEndpointsProtocol=http;BlobEndpoint=http://127.0.0.1:10000/devstoreaccount1;QueueEndpoint=http://127.0.0.1:10001/devstoreaccount1;TableEndpoint=http://127.0.0.1:10002/devstoreaccount1;"

# Initialize BlobServiceClient
blob_service_client = BlobServiceClient.from_connection_string(azure_storage_connection_string)

# Set your OpenAI API key
api_key = "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"
openai.api_key = api_key

app = Flask(__name__)
CORS(app, origins=["http://localhost:3000"])  # Allow requests from React app

# Function to read text from different types of documents
def read_text_from_document(blob_client, extension):
    try:
        if extension == '.docx':
            doc = docx.Document(io.BytesIO(blob_client.download_blob().readall()))
            return '\n'.join([paragraph.text for paragraph in doc.paragraphs])
        elif extension == '.xlsx':
            excel_data = pd.read_excel(io.BytesIO(blob_client.download_blob().readall()))
            return excel_data.to_string(index=False)
        elif extension == '.pptx':
            presentation = pptx.Presentation(io.BytesIO(blob_client.download_blob().readall()))
            return '\n'.join([slide.text for slide in presentation.slides])
        elif extension == '.pdf':
            pdf_reader = PyPDF2.PdfFileReader(io.BytesIO(blob_client.download_blob().readall()))
            return '\n'.join([pdf_reader.getPage(page).extractText() for page in range(pdf_reader.numPages)])
        # elif extension in ['.odt', '.ods', '.odp']:
        #     document = odf.opendocument.load(io.BytesIO(blob_client.download_blob().readall()))
        #     return ' '.join([element.text for element in document.getElementsByType(odf.text.P)]).strip()
        else:
            return blob_client.download_blob().content_as_text()
    except Exception as ex:
        logging.error(f"Error reading document: {str(ex)}")
        return None

# Function to extract text from blobs and concatenate into a single variable
def extract_text_from_blobs(container_name):
    concatenated_text = ""
    container_client = blob_service_client.get_container_client(container_name)

    for blob in container_client.list_blobs():
        blob_name = blob.name
        blob_client = container_client.get_blob_client(blob_name)
        extension = blob_name.lower()[-5:]

        text = read_text_from_document(blob_client, extension)
        if text:
            concatenated_text += text + "\n"

    return concatenated_text.strip()

@app.route('/context_query', methods=['POST'])
def context_query():
    try:
        # Retrieve JSON data from the request
        data = request.get_json()

        # Check if 'user_query' is present in the JSON data
        if 'user_query' not in data:
            return jsonify({"error": "Missing 'user_query' field in the request."}), 400

        # Access 'user_query' directly
        user_query = data['user_query']

        container_name = "uploaded-documents"
        concatenated_text = extract_text_from_blobs(container_name)

        # Call OpenAI to process the text with the user query
        response = openai.ChatCompletion.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "Note 1: If user says Hi, respond with greeting: 'Hello how may I help you ?'.Note2: First Understand the user content and read your content word by word then give the response.Note3: Some data in the content will not be in structured manner you should read that type of content also.Avoid blank space in the content.If the information is not found, respond with 'No Content Found'"},
                {"role": "user", "content": f"{user_query}, Search in the Content and give the response if the content is not found simply return No Content"},
                {"role": "assistant", "content": concatenated_text}
            ]
        )

        return jsonify({"response": response.choices[0].message['content']})
        
    except Exception as e:
        # Log the error for debugging
        error_message = str(e)
        logging.error(f"Error: {error_message}")
        return jsonify({"error": "An internal server error occurred."}), 500

if __name__ == "__main__":
    app.run(debug=True, port=5001)
